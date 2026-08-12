#!/usr/bin/env python3
"""Generate a plain form-porting review workbook from inventory and crosswalk JSON.

Usage:
  python3 build-workbook.py --inventory inventory.json [--crosswalk crosswalk-local.json]
      [--metadata-report metadata-report.json] [--out workbook.xlsx]
      [--mode Review|Exercise|"Release candidate"] [--target-implementation NAME]
      [--target-base-url URL] [--source-form NAME]

Inputs come from extract-inventory.mjs, crosswalk-local.mjs, and fetch-target-metadata.mjs.
The workbook is generated from scratch: no template, no styling, no formulas, no row
limit. Works on Python 3.9+ and needs only openpyxl.
"""

import argparse
import json
import sys

try:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit(
        "openpyxl is required. Install it in an isolated virtual environment, for example:\n"
        "  python3 -m venv .venv && .venv/bin/pip install openpyxl && .venv/bin/python scripts/build-workbook.py ..."
    )


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inventory", required=True)
    parser.add_argument("--crosswalk")
    parser.add_argument("--metadata-report", dest="metadata_report")
    parser.add_argument("--out", default="form-porting-review.xlsx")
    parser.add_argument("--mode", default="Review", choices=["Review", "Exercise", "Release candidate"])
    parser.add_argument("--target-implementation", dest="target_implementation", default="AMRS")
    parser.add_argument("--target-base-url", dest="target_base_url", default="https://kibana.ampath.or.ke/openmrs")
    parser.add_argument("--source-form", dest="source_form")
    parser.add_argument("--open", action="store_true", dest="open_after",
                        help="Open the workbook in the default spreadsheet app after writing it.")
    return parser.parse_args()


def load_json(path):
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def concept_verdict(concept, local, live, has_metadata):
    """(match status, evidence) — terse and decision-relevant. Only live target
    metadata is authoritative: local definitions or usage in other forms do not
    prove a concept exists on the target."""
    if not concept.get("uuid"):
        return "No safe live match", "No concept in source"
    definitions = local.get("definitions") or []
    defined = "; defined in " + definitions[0]["file"].rsplit("/", 1)[-1] if definitions else ""
    if not has_metadata or not live:
        return None, "Not checked live" + defined
    if not live.get("exists"):
        return "No safe live match", "Not on target" + defined
    if live.get("retired"):
        # Exists but retired is a different decision (unretire vs replace) than missing.
        return "Needs team confirmation", "Retired on target" + defined
    return "Strong live match", "Active on target" + defined


def append_rows(sheet, headers, rows):
    sheet.append(headers)
    for row in rows:
        sheet.append(row)


def main():
    args = parse_args()
    inventory = load_json(args.inventory)
    crosswalk = load_json(args.crosswalk) if args.crosswalk else {}
    metadata = load_json(args.metadata_report) if args.metadata_report else {}
    crosswalk_concepts = {c["uuid"]: c for c in crosswalk.get("concepts", [])}
    metadata_concepts = metadata.get("concepts", {})

    form = inventory.get("form", {})
    questions = inventory.get("questions", [])
    concepts = inventory.get("concepts", [])
    checked_on = metadata.get("checkedAt") or crosswalk.get("checkedAt") or ""
    checked_on = checked_on[:10]

    workbook = openpyxl.Workbook()

    decisions = workbook.active
    decisions.title = "Concept Decisions"
    rows = []
    for index, concept in enumerate(concepts):
        uuid = concept["uuid"]
        local = crosswalk_concepts.get(uuid, {})
        live = metadata_concepts.get(uuid, {})
        definitions = local.get("definitions") or []
        proposed = live.get("name") or (definitions[0]["name"] if definitions else None)
        status, evidence = concept_verdict(concept, local, live, bool(metadata_concepts))
        rows.append((
            index + 1,
            ", ".join(concept.get("roles", [])),
            "; ".join(concept.get("labels", [])) or uuid,
            proposed,
            uuid or None,
            status,
            evidence,
            None,  # Team decision
            None,  # Approved target UUID
            checked_on or None,
            None,  # Reviewer notes
        ))
    append_rows(decisions, (
        "#", "Source type", "Source item", "Proposed target concept", "Concept UUID",
        "Match status", "Evidence", "Team decision", "Approved target UUID", "Checked on", "Reviewer notes",
    ), rows)
    # apply-decisions.py matches the decision phrases literally, so constrain the column.
    dv = DataValidation(
        type="list",
        formula1='"Approve target concept,Create new concept,Redesign question,Remove from target form,Needs team decision"',
        allow_blank=True, showErrorMessage=True,
    )
    decisions.add_data_validation(dv)
    dv.add("H2:H{}".format(len(rows) + 1))
    decisions["H1"].comment = Comment(
        "Pick one of the five options from the dropdown. Fill Approved target UUID when approving or mapping. "
        "Do not rename column headers - the apply script finds columns by name.", "port-openmrs-form")

    design = workbook.create_sheet("Form Design")
    rows = []
    for index, question in enumerate(questions):
        expressions = question.get("expressions", {})
        messages = [v.get("message") for v in question.get("validators", []) if isinstance(v, dict) and v.get("message")]
        rows.append((
            question.get("order", index + 1),
            question.get("page"),
            question.get("section"),
            question.get("label"),
            question.get("id"),
            question.get("type"),
            question.get("rendering"),
            question.get("concept"),
            question.get("required"),
            expressions.get("hideWhenExpression"),
            "; ".join(messages) if messages else None,
            None,  # Team decision
            None,  # Notes
        ))
    append_rows(design, (
        "Order", "Page", "Section", "Label", "Field ID", "Type", "Rendering", "Concept UUID",
        "Required", "Show or hide rule", "Validation message", "Team decision", "Notes",
    ), rows)

    results = workbook.create_sheet("Build Results")
    # The "Encounter type UUID" cell is what apply-decisions.py treats as decided, so it
    # only carries the source value when live metadata confirms it exists and is active.
    live_encounter = metadata.get("encounterType") or {}
    declared_encounter = form.get("encounterType")
    encounter_confirmed = (bool(declared_encounter) and live_encounter.get("uuid") == declared_encounter
                           and live_encounter.get("exists") and not live_encounter.get("retired"))
    entries = [
        ("Form name", form.get("name")),
        ("Version", form.get("version")),
        ("Mode", args.mode),
        ("Target implementation", args.target_implementation),
        ("Target base URL", args.target_base_url),
        ("Encounter type name", live_encounter.get("name") if encounter_confirmed else None),
        ("Encounter type UUID", declared_encounter if encounter_confirmed else None),
        ("Form UUID", form.get("uuid")),
        ("Source form", args.source_form or form.get("path")),
        ("Checked on", checked_on),
    ]
    if declared_encounter and not encounter_confirmed:
        entries.append(("Source encounter type UUID (not verified on target)", declared_encounter))
    for key, value in entries:
        results.append((key, value))
    results.append(())
    append_rows(results, ("Step", "Gate", "Status", "Evidence", "Next action", "Artifact"), [
        (1, "Source captured", "Not started", None, None, None),
        (2, "Target crosswalk complete", "Not started", None, None, None),
        (3, "Team decisions complete", "Not started", None, None, None),
        (4, "Design and encounter metadata complete", "Not started", None, None, None),
        (5, "O3 JSON generated", "Not started", None, None, None),
        (6, "Canonical and semantic validation passed", "Not started", None, None, None),
        (7, "No-save preview passed", "Not started", None, None, None),
    ])
    results.append(())
    results.append(("Notes", "Team fills only: Team decision (dropdown), Approved target UUID, and Reviewer notes "
                             "on Concept Decisions; the Encounter type and Form UUID cells above. Do not rename headers."))

    workbook.save(args.out)
    print("Wrote {} ({} concept rows, {} design rows).".format(args.out, len(concepts), len(questions)))
    print("Match status is proposed from live metadata only (blank without it); Team decision columns are left for review.")
    if args.open_after:
        import platform
        import subprocess
        try:
            system = platform.system()
            if system == "Darwin":
                subprocess.run(["open", args.out], check=False)
            elif system == "Windows":
                import os
                os.startfile(args.out)  # noqa: S606
            else:
                subprocess.run(["xdg-open", args.out], check=False)
        except Exception as error:
            print("Could not open the workbook automatically: {}".format(error))


if __name__ == "__main__":
    main()
